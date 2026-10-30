from pathlib import Path
from html.parser import HTMLParser
from zipfile import ZipFile
from docx import Document
from pptx import Presentation
import fitz
ROOT=Path(__file__).resolve().parent
for lang in ['AR','EN']:
 for stem,ext,count in [('PMP-Lean-Operating-Guide','docx',5),('PMP-Lean-Playbook','pptx',10)]:
  f=ROOT/f'{stem}-{lang}.{ext}'
  with ZipFile(f) as z:assert z.testzip() is None
  pdf=fitz.open(f.with_suffix('.pdf'));assert len(pdf)==count
  assert all(len(p.get_text().strip())>500 for p in pdf)
  if ext=='pptx':
   deck=Presentation(f);assert len(deck.slides)==10
   assert not any(sh.shape_type==13 for sl in deck.slides for sh in sl.shapes)
   for i,page in enumerate(pdf):
    for b in page.get_text('blocks'):
     assert b[1]>=0 and b[3]<=page.rect.height
     if 240<b[1]<420:assert b[3]<434,(lang,i+1,'card overflow',b[:4])
  else:
   d=Document(f);t=' '.join(p.text for p in d.paragraphs)
   for term in ['VE-01','VE-02','GD-01','GD-02','Claude','ChatGPT','Canva','Adobe','Higgsfield','SMART']:assert term in t,term
   if lang=='AR':
    assert 'المهمة تفضل مفتوحة' in t
    assert any(p._p.xpath('.//w:bCs') for p in d.paragraphs),'Arabic bold missing'
  print(f.name,count,'pages/slides: PASS')
class Links(HTMLParser):
 def __init__(self,base):super().__init__();self.base=base
 def handle_starttag(self,tag,attrs):
  if tag in ['a','iframe']:
   href=dict(attrs).get('href' if tag=='a' else 'src');assert href and (self.base/href).is_file(),href
for folder in ['lean','arabic']:
 p=ROOT.parent/folder;Links(p).feed((p/'index.html').read_text())
print('All bilingual download and embedded preview links exist.')
