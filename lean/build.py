from pathlib import Path
from content import SECTIONS
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from docx import Document
from docx.shared import Inches as DI,Pt as DP,RGBColor as DC
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
OUT=Path(__file__).parent
BURG='7D2B42';YELLOW='FFCC00';INK='202020';MUTED='555555';PAPER='F5F3F0';WHITE='FFFFFF'
def color(v):return RGBColor.from_string(v)
def box(sl,x,y,w,h,fill):
 s=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));s.fill.solid();s.fill.fore_color.rgb=color(fill);s.line.fill.background();return s
def txt(sl,x,y,w,h,text,size=18,bold=False,ink=INK,ar=False):
 s=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=s.text_frame;tf.word_wrap=True
 tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if not i else tf.add_paragraph();p.alignment=PP_ALIGN.RIGHT if ar else PP_ALIGN.LEFT;p._p.get_or_add_pPr().set('rtl','1' if ar else '0');p.line_spacing=1.12;p.space_after=Pt(4)
  r=p.add_run();r.text=line;r.font.name='IBM Plex Sans Arabic' if ar else 'Montserrat';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=color(ink)
 return s
for lang in ['en','ar']:
 ar=lang=='ar';suffix=lang.upper();deck=Presentation();deck.slide_width=Inches(13.333);deck.slide_height=Inches(7.5)
 for i,sec in enumerate(SECTIONS):
  title,sub,blocks,gate,record=sec[lang];sl=deck.slides.add_slide(deck.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=color(PAPER)
  box(sl,0,0,13.333,.13,BURG)
  txt(sl,.55,.31,12.2,.35,'PMP  /  دليل التشغيل المختصر' if ar else 'PMP  /  THE LEAN OPERATING PLAYBOOK',12,True,BURG,ar)
  txt(sl,.55,.91,12.2,.68,title,30,True,BURG,ar)
  txt(sl,.55,1.68,12.2,.55,sub,16,False,MUTED,ar)
  for j,(heading,body) in enumerate(blocks):
   x=.55+(2-j if ar else j)*4.12
   box(sl,x,2.47,3.99,3.55,WHITE);box(sl,x,2.47,3.99,.05,YELLOW)
   txt(sl,x+.19,2.7,3.61,.66,heading,18,True,BURG,ar)
   txt(sl,x+.19,3.47,3.61,2.36,body,15.5 if ar else 15,False,INK,ar)
  box(sl,.55,6.2,12.23,.75,BURG)
  txt(sl,.77,6.3,11.79,.6,gate,15 if ar else 14.5,True,WHITE,ar)
  txt(sl,.55,7.14,11.5,.2,'قواعد مؤكدة؛ المقترحات والنواقص مكتوبة بوضوح.' if ar else 'Confirmed rules; proposals and open decisions are marked.',9,False,MUTED,ar)
  txt(sl,12.17,7.08,.6,.3,f'{i+1:02d}',12,True,BURG)
  sl.notes_slide.notes_text_frame.text=record
 deck.save(OUT/f'PMP-Lean-Playbook-{suffix}.pptx')
 d=Document();s=d.sections[0];s.page_height=DI(11.7);s.page_width=DI(8.3);s.top_margin=DI(.65);s.bottom_margin=DI(.6);s.left_margin=s.right_margin=DI(.7)
 normal=d.styles['Normal'];normal.font.name='IBM Plex Sans Arabic' if ar else 'Montserrat';normal.font.size=DP(12.5 if ar else 12);normal.paragraph_format.line_spacing=1.14;normal.paragraph_format.space_after=DP(8)
 for name,sz in [('Title',24),('Heading 1',19),('Heading 2',12)]:
  st=d.styles[name];st.font.name=normal.font.name;st.font.size=DP(sz);st.font.bold=True;st.font.color.rgb=DC.from_string(BURG);st.paragraph_format.space_before=DP(8);st.paragraph_format.space_after=DP(8)
 def para(text,style=None,bold=False,shade=None):
  p=d.add_paragraph(text,style);p.alignment=WD_ALIGN_PARAGRAPH.RIGHT if ar else WD_ALIGN_PARAGRAPH.LEFT
  if ar:
   b=OxmlElement('w:bidi');p._p.get_or_add_pPr().append(b)
  for r in p.runs:
   r.font.name=normal.font.name
   if bold:r.bold=True
   rp=r._r.get_or_add_rPr();fonts=rp.rFonts
   if fonts is not None:fonts.set(qn('w:cs'),normal.font.name)
  if shade:
   el=OxmlElement('w:shd');el.set(qn('w:fill'),shade);p._p.get_or_add_pPr().append(el)
  return p
 groups=[[0,1],[2,3],[4,5],[6,7],[8,9]]
 md=[]
 for page,group in enumerate(groups):
  if page:d.add_page_break()
  for idx in group:
   title,sub,blocks,gate,record=SECTIONS[idx][lang]
   para(f'{idx+1:02d}  /  {title}','Heading 1');para(sub,bold=True)
   md+=['## '+title,sub]
   for label,body in blocks:
    p=para('');r=p.add_run(label+' — ');r.bold=True;p.add_run(body)
    for run in p.runs:run.font.name=normal.font.name
    md += ['**'+label+'** — '+body]
   para(gate,bold=True,shade='F4E8EC');md+=[gate]
   p=para(record);md+=[record]
   for run in p.runs:run.font.size=DP(9.5);run.font.color.rgb=DC.from_string(MUTED)
 header=s.header.paragraphs[0];header.text='PMP  |  دليل تشغيل عملي · نسخة مختصرة' if ar else 'PMP  |  Practical operating guide · Lean edition';header.style=d.styles['Caption'];header.alignment=WD_ALIGN_PARAGRAPH.RIGHT if ar else WD_ALIGN_PARAGRAPH.LEFT
 footer=s.footer.paragraphs[0];footer.text='٣٠ أكتوبر ٢٠٢٦  •  للمالك ومدير التسويق  |  ' if ar else '30 October 2026  •  Owner & Marketing Manager  |  ';footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT if ar else WD_ALIGN_PARAGRAPH.LEFT
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
 for p in [header,footer]:
  for r in p.runs:r.font.name=normal.font.name;r.font.size=DP(8);r.font.color.rgb=DC.from_string(MUTED)
 if ar:
  for p in d.paragraphs:
   for run in p.runs:
    rp=run._r.get_or_add_rPr()
    rf=rp.rFonts
    if rf is not None:rf.set(qn('w:cs'),'IBM Plex Sans Arabic')
    sz=run.font.size or p.style.font.size or DP(12.5)
    cs=OxmlElement('w:szCs');cs.set(qn('w:val'),str(int(sz.pt*2)));rp.append(cs)
    if run.bold or (run.bold is None and p.style.font.bold):
     b=OxmlElement('w:bCs');rp.append(b)
 d.save(OUT/f'PMP-Lean-Operating-Guide-{suffix}.docx');(OUT/f'PMP-Lean-Operating-Guide-{suffix}.md').write_text('\n\n'.join(md))
 print(suffix,'guide and 10-slide deck built')
