from pathlib import Path
from zipfile import ZipFile
from docx import Document
from pptx import Presentation
import json
from html.parser import HTMLParser
root=Path(__file__).resolve().parent
for name in ['PMP-Incoming-Marketing-Manager-Detailed-Runbook.docx','PMP-Incoming-Marketing-Manager-Detailed-Onboarding.pptx']:
 with ZipFile(root/name) as z:assert z.testzip() is None
word=Document(root/'PMP-Incoming-Marketing-Manager-Detailed-Runbook.docx')
text='\n'.join(p.text for p in word.paragraphs)+'\n'+'\n'.join(c.text for t in word.tables for row in t.rows for c in row.cells)
md=(root/'PMP-Incoming-Marketing-Manager-Detailed-Runbook.md').read_text()
for i in range(1,15):assert f'SOP-{i:02d}' in text
for role in ['MM-01','BD-01','SA-01','AM-01','HC-01','CW-S01','CW-J01','HD-01','GD-01','GD-02','VE-01','VE-02','SM-01','MB-01']:assert role in text
for term in ['Higgsfield','Claude','ChatGPT','Canva','Adobe','hard drives','memory cards','studio tools','behind-the-scenes','podcast','PR person','readiness to sign','100%','50%','25%','NOT ASSESSED']:
 assert term in text,term
 assert term in md,term
assert 'all other team members are forbidden' in text.lower()
assert 'including MM' in text
assert 'do not assume SM automatically owns every hard drive' not in text or 'custodian' in text
for heading in ['Company subscriptions','Shoot logistics','Review quality','Daily operating','Practice scenarios','Owner decisions']:assert heading.lower() in text.lower(),heading
content=json.loads((root/'onboarding-content.json').read_text());assert len(content['managerChecks'])==14;assert len(content['templates'])==15
assert len(content['practiceScenarios'])==17
ppt=Presentation(root/'PMP-Incoming-Marketing-Manager-Detailed-Onboarding.pptx');assert len(ppt.slides)==51
assert not any(sh.shape_type==13 for sl in ppt.slides for sh in sl.shapes)
ppttext='\n'.join(sh.text for sl in ppt.slides for sh in sl.shapes if sh.has_text_frame)
assert 'VE-01/VE-02 ONLY' in ppttext and 'Forbidden for every other team member' in ppttext
for term in ['memory cards','Higgsfield','hard drives','PR person']:assert term.lower() in ppttext.lower()
class Links(HTMLParser):
 def handle_starttag(self,tag,attrs):
  if tag=='a':
   link=dict(attrs).get('href');assert link and (root/link).exists(),link
Links().feed((root/'index.html').read_text())
import fitz
for name in ['PMP-Incoming-Marketing-Manager-Detailed-Runbook','PMP-Incoming-Marketing-Manager-Detailed-Onboarding']:
 p=fitz.open(root/(name+'.pdf'));assert len(p)>0
 for i,page in enumerate(p):
  assert len(page.get_text().strip())>120,(name,i+1,'unexpected near-blank page')
  for block in page.get_text('blocks'):
   assert block[1]>=0 and block[3]<=page.rect.height+1,(name,i+1,'text beyond page')
 print(name+': '+str(len(p))+' rendered pages')
print('PASS: 14 procedures; 14 distinct roles; confirmed resource/custody rules; 17 practice scenarios; editable 51-slide training; intact Office packages; complete Markdown; local download links.')
