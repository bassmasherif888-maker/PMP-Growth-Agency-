from pathlib import Path
from zipfile import ZipFile
from docx import Document
from pptx import Presentation
import json
root=Path(__file__).resolve().parent
for name in ['PMP-Marketing-Department-Owner-Handover-SOP.docx','PMP-Marketing-Department-Owner-Handover.pptx']:
 with ZipFile(root/name) as z:assert z.testzip() is None
sop=Document(root/'PMP-Marketing-Department-Owner-Handover-SOP.docx')
text='\n'.join(p.text for p in sop.paragraphs)+'\n'+'\n'.join(c.text for t in sop.tables for row in t.rows for c in row.cells)
register=json.loads((root/'procedure-register.json').read_text());assert len(register)==12
for p in register:
 assert p['id']+' | '+p['title'] in text
 for key in ['lead','trigger','inputs','steps','output','gate','record']:assert p[key]
for id in ['MM-01','BD-01','SA-01','AM-01','HC-01','CW-S01','CW-J01','HD-01','GD-01','GD-02','VE-01','VE-02','SM-01','MB-01']:assert id in text
for phrase in ['readiness to sign','100%','50%','25%','Daily operating rhythm','owner','NOT ASSESSED','same calendar','pending']:
 assert phrase.lower() in text.lower(),phrase
md=(root/'PMP-Marketing-Department-Owner-Handover-SOP.md').read_text()
for heading in ['Daily operating rhythm','Tools and records','Approval matrix','Sources and limits']:assert heading in md
p=Presentation(root/'PMP-Marketing-Department-Owner-Handover.pptx');assert len(p.slides)==23
assert not any(sh.shape_type==13 for sl in p.slides for sh in sl.shapes)
from html.parser import HTMLParser
class Links(HTMLParser):
 def handle_starttag(self,tag,attrs):
  if tag=='a':
   href=dict(attrs).get('href');assert href and ((root/href).is_file() or ((root/href).is_dir() and (root/href/'index.html').is_file())),href
Links().feed((root/'index.html').read_text())
try:
 import fitz
 for stem in ['PMP-Marketing-Department-Owner-Handover-SOP','PMP-Marketing-Department-Owner-Handover']:
  pdf=fitz.open(root/(stem+'.pdf'))
  assert len(pdf)>0
  for page in pdf:
   assert len(page.get_text().strip())>100,'Unexpected blank page'
  print(stem+': '+str(len(pdf))+' rendered pages')
except ImportError:print('PDF text check skipped: install PyMuPDF for rendering checks.')
print('Validated: 12 complete procedures, 14 separate positions, editable Office files, complete Markdown, 23-slide briefing and working landing-page file links.')
